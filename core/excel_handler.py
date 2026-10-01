# Copyright (c) 2026 The Forever CSF
# SPDX-License-Identifier: AGPL-3.0-or-later
# Additional terms: see 第二部分《标准处理系统 许可、隐私与免责说明》（AGPL-3.0 第 7 条附加条款）
import os
import shutil
import tempfile
import threading
from typing import Dict, Any, List, Callable

from utils.helpers import search_standard_names, fill_merged_cells_down, resolve_standard_display

# Excel 的 COM 自动化同样是单实例的：多线程同时 Dispatch/Open/Quit 会互相顶掉。
# 只锁 COM 转换这一段，其它文件类型仍并行。
_EXCEL_COM_LOCK = threading.RLock()


class ExcelHandler:
    """Excel 处理管道：合并单元格识别、向下/向右填充、公式转值、标准名称识别。"""

    def __init__(self, on_log: Callable[[str], None] = None,
                 on_progress: Callable[[int], None] = None):
        self.on_log = on_log or (lambda x: None)
        self.on_progress = on_progress or (lambda x: None)
        self.cancel_check = lambda: None

    # ── 旧版 .xls 转换 ────────────────────────────────────
    @staticmethod
    def is_legacy_xls(file_path: str) -> bool:
        return os.path.splitext(file_path)[1].lower() == ".xls"

    @staticmethod
    def convert_xls_to_xlsx(xls_path: str, on_log: Callable[[str], None] = None) -> str:
        """把旧版 .xls 另存为 .xlsx。

        为什么要转：本程序用 openpyxl 读表格，而 openpyxl **只能读 xlsx**，
        遇到 .xls 会直接抛异常（用户看到的是一句看不懂的报错）。
        这里借本机 Excel 的 COM 接口转换，和 .doc → .docx 是同一套思路。

        注意：
        * 用 DispatchEx 强制开新实例，避免连上用户自己开着的 Excel（随后 Quit
          会把用户没保存的工作簿一起关掉）；
        * DisplayAlerts=False + 只读打开，避免弹框把线程挂死；
        * 转换结果放到临时目录，不碰用户原文件。
        """
        log = on_log or (lambda x: None)
        out_dir = None
        excel = None
        wb = None
        com_inited = False
        with _EXCEL_COM_LOCK:
            try:
                try:
                    import pythoncom
                    pythoncom.CoInitialize()
                    com_inited = True
                except Exception:
                    com_inited = False
                import win32com.client as win32
                try:
                    excel = win32.DispatchEx("Excel.Application")
                except Exception:
                    excel = win32.Dispatch("Excel.Application")
                try:
                    excel.Visible = False
                    excel.DisplayAlerts = False
                except Exception:
                    pass
                out_dir = tempfile.mkdtemp()
                out_path = os.path.join(
                    out_dir, os.path.splitext(os.path.basename(xls_path))[0] + ".xlsx")
                wb = excel.Workbooks.Open(os.path.abspath(xls_path),
                                          ReadOnly=True, UpdateLinks=0)
                # 51 = xlOpenXMLWorkbook (.xlsx)
                wb.SaveAs(out_path, FileFormat=51)
                wb.Close(False)
                wb = None
                excel.Quit()
                excel = None

                # 落到持久目录，避免 mkdtemp 被清理后路径失效
                persistent_dir = os.path.join(tempfile.gettempdir(),
                                              "standard_system", "_xls_converted")
                os.makedirs(persistent_dir, exist_ok=True)
                final_path = os.path.join(persistent_dir, os.path.basename(out_path))
                shutil.copy2(out_path, final_path)
                log(f"  [转换完成] {os.path.basename(xls_path)} → .xlsx")
                return final_path
            except Exception as e:
                log(f"  [转换失败] 旧版 .xls 转换失败（{e}）。"
                    f"请用 Excel 另存为 .xlsx 后重试。")
                raise
            finally:
                try:
                    if wb is not None:
                        wb.Close(False)
                except Exception:
                    pass
                try:
                    if excel is not None:
                        excel.Quit()
                except Exception:
                    pass
                if out_dir and os.path.exists(out_dir):
                    shutil.rmtree(out_dir, ignore_errors=True)
                if com_inited:
                    try:
                        import pythoncom
                        pythoncom.CoUninitialize()
                    except Exception:
                        pass

    def process(self, file_path: str, options: Dict[str, Any]) -> Dict[str, Any]:
        import openpyxl
        self.on_log(f"  [处理Excel] {file_path}")

        wb_data = openpyxl.load_workbook(file_path, data_only=True)
        wb_formula = openpyxl.load_workbook(file_path, data_only=False)
        result = {"tables": [], "text": "", "logs": [], "toc": []}
        sheets = wb_data.sheetnames

        for s_idx, sheet_name in enumerate(sheets):
            self.cancel_check()   # 支持"终止处理"即时生效
            ws_data = wb_data[sheet_name]
            ws_formula = wb_formula[sheet_name]
            self.on_log(f"  处理工作表: {sheet_name}")
            merged_map = self._build_merged_map(ws_data)
            rows = []
            for row_idx, row in enumerate(ws_data.iter_rows(values_only=False), start=1):
                if row_idx % 200 == 0:
                    self.cancel_check()
                row_data = []
                for col_idx, cell in enumerate(row, start=1):
                    formula_cell = ws_formula.cell(row_idx, col_idx)
                    value = self._get_cell_value(cell, formula_cell, merged_map, row_idx, col_idx)
                    row_data.append(value)
                rows.append(row_data)

            # 合并单元格填充
            rows = self._fill_merged(rows, merged_map)
            rows = fill_merged_cells_down(rows)

            for r_idx, cells in enumerate(rows, start=1):
                if r_idx % 200 == 0:
                    self.cancel_check()
                content = " | ".join(str(c) for c in cells)
                for name in search_standard_names(content):
                    display_name = resolve_standard_display(content, name)
                    result["tables"].append({
                        "name": display_name,
                        "content": content,
                        "page": f"{sheet_name} 行{r_idx}",
                    })
            result["text"] += f"\n[工作表: {sheet_name}]\n" + "\n".join(
                " | ".join(str(c) for c in row) for row in rows
            )
            self.on_progress(int((s_idx + 1) / len(sheets) * 100 + 0.5))

        wb_data.close()
        wb_formula.close()
        return result

    def preview(self, file_path: str) -> Dict[str, Any]:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        merged_map = self._build_merged_map(ws)
        rows = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
            row_data = []
            for col_idx, cell in enumerate(row, start=1):
                value = self._get_cell_value(cell, None, merged_map, row_idx, col_idx)
                row_data.append(value)
            rows.append(row_data)
        rows = self._fill_merged(rows, merged_map)
        rows = fill_merged_cells_down(rows)

        # 生成 HTML 表格
        html_parts = ['<table style="width:100%;border-collapse:collapse;border:1px solid #CBD5E1;font-size:12px;">']
        for r_idx, row in enumerate(rows[:50], start=1):
            html_parts.append("<tr>")
            for cell in row:
                text = str(cell) if cell is not None else ""
                text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                html_parts.append(f'<td style="border:1px solid #CBD5E1;padding:4px 6px;">{text}</td>')
            html_parts.append("</tr>")
        html_parts.append("</table>")

        lines = ["合并单元格已填充："] + [" | ".join(str(c) for c in row) for row in rows[:50]]
        return {
            "html": "\n".join(html_parts),
            "text": "\n".join(lines),
            "columns": rows[0] if rows else [],
            "rows": rows[1:51] if len(rows) > 1 else [],
        }

    def _build_merged_map(self, ws):
        merged_map = {}
        for merged_range in ws.merged_cells.ranges:
            min_row, min_col = merged_range.min_row, merged_range.min_col
            max_row, max_col = merged_range.max_row, merged_range.max_col
            value = ws.cell(min_row, min_col).value
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    merged_map[(r, c)] = (min_row, min_col, value)
        return merged_map

    def _get_cell_value(self, cell, formula_cell, merged_map, row_idx, col_idx):
        if (row_idx, col_idx) in merged_map:
            _, _, value = merged_map[(row_idx, col_idx)]
            return value if value is not None else ""
        value = cell.value
        if value is None and formula_cell is not None and formula_cell.value is not None:
            fv = formula_cell.value
            if isinstance(fv, str) and fv.startswith("="):
                return fv
        return value if value is not None else ""

    def _fill_merged(self, rows: List[List[Any]], merged_map: Dict) -> List[List[Any]]:
        """根据合并单元格映射向右填充（跨列）。"""
        if not rows or not merged_map:
            return rows
        filled = [list(row) for row in rows]
        col_count = max(len(r) for r in filled)
        for row in filled:
            while len(row) < col_count:
                row.append("")
        for (row_idx, col_idx), (min_row, min_col, value) in merged_map.items():
            r = row_idx - 1
            c = col_idx - 1
            if 0 <= r < len(filled) and 0 <= c < col_count:
                filled[r][c] = value if value is not None else ""
        return filled
